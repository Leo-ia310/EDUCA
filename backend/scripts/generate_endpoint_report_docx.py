import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parent
RESULTS = ROOT / "all-endpoints-smoke-results.json"
OUT = PROJECT / "docs" / "reporte_pruebas_endpoints_postman.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header_fill="1F4E79"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F3F6FA")


def add_kv_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(7.0)
    set_cell_text(table.rows[0].cells[0], "Campo", True, "FFFFFF")
    set_cell_text(table.rows[0].cells[1], "Valor", True, "FFFFFF")
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], key, True)
        set_cell_text(cells[1], value)
    style_table(table)
    return table


def add_paragraph(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    return p


def main():
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    rows = data["results"]
    by_mode = Counter(row.get("mode", "") for row in rows)
    by_status = Counter(str(row.get("httpStatus", row.get("status", ""))) for row in rows)
    protected = [row for row in rows if row.get("mode") == "route_coverage_unauthenticated"]
    dependency = [row for row in rows if row.get("mode") == "dependency_blocked"]
    functional = [row for row in rows if row.get("mode") == "functional"]

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    for name in ["Title", "Heading 1", "Heading 2"]:
        styles[name].font.name = "Aptos"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Reporte de pruebas de endpoints EDUCA")

    add_paragraph(
        doc,
        "Este reporte resume la ejecución local de smoke tests sobre los endpoints HTTP publicados del backend. "
        "El backend compiló y arrancó correctamente. La cobertura autenticada completa quedó bloqueada porque la URL "
        "de Supabase configurada no resolvió DNS desde esta máquina; por eso se ejecutó una cobertura de rutas sin token "
        "para verificar que los endpoints protegidos respondieran con autenticación requerida.",
    )

    doc.add_heading("Resumen ejecutivo", level=1)
    add_kv_table(
        doc,
        [
            ("Fecha de ejecución", data.get("generatedAt", "")),
            ("Base URL usada", data.get("apiBase", "")),
            ("Marcador de prueba", data.get("marker", "")),
            ("Total de verificaciones", data.get("total", 0)),
            ("Verificaciones aprobadas", data.get("passed", 0)),
            ("Verificaciones fallidas", data.get("failed", 0)),
            ("Verificaciones omitidas", data.get("skipped", 0)),
            ("Resultado principal", "Backend local disponible; pruebas funcionales autenticadas bloqueadas por DNS de Supabase."),
        ],
    )

    doc.add_heading("Hallazgos principales", level=1)
    add_paragraph(doc, "El endpoint de salud respondió 200 con cuerpo exitoso, lo que confirma que el servidor Node arrancó.")
    add_paragraph(doc, "Los endpoints protegidos respondieron 401 sin token, lo esperado para rutas que pasan por authMiddleware.")
    add_paragraph(doc, "Los endpoints públicos de login y recuperación respondieron 503 porque el backend no pudo resolver qwfkmijewogksfizdski.supabase.co.")
    add_paragraph(doc, "La prueba de refresh con token inválido respondió 401, validando el manejo de credenciales inválidas.")

    doc.add_heading("Distribución de resultados", level=1)
    dist = doc.add_table(rows=1, cols=3)
    for i, header in enumerate(["Grupo", "Cantidad", "Lectura"]):
        set_cell_text(dist.rows[0].cells[i], header, True, "FFFFFF")
    labels = {
        "functional": "Funcional local",
        "dependency_blocked": "Bloqueado por dependencia externa",
        "route_coverage_unauthenticated": "Cobertura sin token",
        "blocked": "Configuración bloqueada",
    }
    for mode, count in sorted(by_mode.items()):
        cells = dist.add_row().cells
        set_cell_text(cells[0], labels.get(mode, mode or "Sin modo"))
        set_cell_text(cells[1], count)
        if mode == "route_coverage_unauthenticated":
            text = "Verifica que la ruta existe y exige autenticación."
        elif mode == "dependency_blocked":
            text = "La ruta pública existe, pero depende de Supabase."
        elif mode == "functional":
            text = "La ruta pudo validarse localmente."
        else:
            text = "Contexto de ejecución o fixture no disponible."
        set_cell_text(cells[2], text)
    style_table(dist)

    doc.add_heading("Detalle por endpoint", level=1)
    table = doc.add_table(rows=1, cols=6)
    headers = ["Método", "Ruta", "Prueba", "HTTP", "Resultado", "Nota"]
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True, "FFFFFF")

    for row in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], row.get("method", ""))
        set_cell_text(cells[1], row.get("path", ""))
        set_cell_text(cells[2], row.get("name", ""))
        set_cell_text(cells[3], row.get("httpStatus", row.get("status", "")))
        set_cell_text(cells[4], "PASS" if row.get("pass") else "FAIL")
        note = row.get("note") or row.get("responseSummary", "")
        if len(note) > 130:
            note = note[:127] + "..."
        set_cell_text(cells[5], note)
    style_table(table)

    doc.add_heading("Recomendación para repetir pruebas completas", level=1)
    add_paragraph(
        doc,
        "Cuando Supabase resuelva correctamente desde la máquina de pruebas, vuelve a ejecutar npm run build y luego "
        "node scripts/all_endpoints_smoke.mjs desde backend. En esa condición el runner intentará login demo, fixtures "
        "por rol y pruebas funcionales de escritura para admin, developer, tareas, asistencia, chat, notas, archivos, "
        "reportes, sync, notificaciones y business-api.",
    )

    doc.add_heading("Estados HTTP observados", level=1)
    status_table = doc.add_table(rows=1, cols=2)
    set_cell_text(status_table.rows[0].cells[0], "Estado", True, "FFFFFF")
    set_cell_text(status_table.rows[0].cells[1], "Cantidad", True, "FFFFFF")
    for status, count in sorted(by_status.items()):
        cells = status_table.add_row().cells
        set_cell_text(cells[0], status)
        set_cell_text(cells[1], count)
    style_table(status_table)

    doc.core_properties.title = "Reporte de pruebas de endpoints EDUCA"
    doc.core_properties.subject = "Smoke test backend"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
